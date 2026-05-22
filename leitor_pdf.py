from __future__ import annotations

import argparse
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import fitz  # PyMuPDF
import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill


OUTPUT_COLUMNS = [
    "Elaborador",
    "Nº do processo",
    "Nº precatorio",
    "Nome do Credor",
    "Entidade/Ente Federado",
    "Data pagamento",
    "Tipo de pagamento",
    "Valor Bruto do Pagamento",
    "Valor líquido pago a parte",
    "NATUREZA",
]

PROCESS_RE = re.compile(r"\b(\d{12})\b")
PAGE_RE = re.compile(r"^\d+/\d+$")
DATE_RE = re.compile(r"\d{2}/\d{2}/\d{4}")
VALUE_RE = re.compile(r"\d{1,3}(?:\.\d{3})*,\d{2}")
PAYMENT_RE = re.compile(
    r"PARCIAL\s+ANTECIPACAO"
    r"|PAGAMENTO\s+ANTECIPACAO"
    r"|PAGAMENTO\s+INTEGRAL"
    r"|CESSAO\s*-\s*ACORDO(?:\s+DIRETO)?"
)
ORPHAN_PAYMENT_RE = re.compile(r"\bPAGAMENTO\b(?!\s+(?:INTEGRAL|ANTECIPACAO))")

# Coordenadas apos remover a rotacao do PDF.
COLUMN_RANGES = {
    "processo": (756.0, 842.0),
    "credor": (627.0, 756.0),
    "entidade": (495.0, 627.0),
    "data_pagamento": (416.0, 495.0),
    "tipo_pagamento": (341.0, 416.0),
    "tipo_antecipacao": (290.0, 341.0),
    "valor_bruto": (216.0, 290.0),
    "previdencia": (182.0, 216.0),
    "imposto_renda": (122.0, 182.0),
    "bloqueio": (69.0, 122.0),
    "valor_liquido": (0.0, 69.0),
}

HEADER_PATTERNS = [
    "TRIBUNAL DE JUSTICA",
    "DEPARTAMENTO DE PRECATORIOS",
    "PRECATORIOS PAGOS NO PERIODO",
    "HORA:",
    "DATA:",
    "N DO PRECATORIO",
    "NOME DO CREDOR",
    "ENTIDADE/ENTE FEDERADO",
    "DATA PAGAMENTO",
    "TIPO DE PAGAMENTO",
    "VALOR BRUTO DO PAGAMENTO",
    "VALOR LIQUIDO PAGO A PARTE",
    "PREVIDENCIA",
    "IMPOSTO DE RENDA",
    "TIPO DE ANTECIPACAO",
    "VALOR DO BLOQUEIO",
]

SUMMARY_LINE_PREFIXES = (
    "TOTAL DE REGISTROS",
    "VALOR TOTAL",
)

HEADER_TOKEN_GUARD = {
    "VALOR",
    "BLOQUEIO",
    "IMPOSTO",
    "RENDA",
    "PREVIDENCIA",
    "ANTECIPACAO",
    "PARTE",
    "LIQUIDO",
}

LINE_Y_TOLERANCE = 1.2
TOP_BAND_PADDING = 12.0
NAME_SEGMENTS_KEY = "__name_segments__"
MANUAL_REVIEW_FLAGS = {
    "multiple_liquid_tokens",
    "multiple_process_ids",
    "liquid_gt_gross",
}
AUTO_ADJUSTMENT_FLAGS = {
    "inferred_payment_type",
    "sequence_adjusted",
}


@dataclass
class ExtractedLine:
    y: float
    columns: dict[str, str]
    text: str
    normalized_text: str


@dataclass
class PaymentComponent:
    payment_type: str
    date: str
    gross_value: float


@dataclass
class PaymentBuildDiagnostics:
    raw_payment_type_count: int
    raw_date_count: int
    raw_gross_count: int
    inferred_payment_type_count: int
    sequence_adjusted: bool


@dataclass
class BuiltRow:
    row: dict[str, object]
    component_count: int
    retained_component_count: int
    had_cession: bool
    warnings: list[str]
    diagnostic_flags: list[str]
    validation_notes: list[str]
    source_page: int


@dataclass
class ExtractionRunResult:
    pdf_files: list[Path]
    built_rows: list[BuiltRow]
    output_path: Path
    report_path: Path


def normalize_ascii(text: str) -> str:
    normalized = unicodedata.normalize("NFKD", text)
    ascii_text = normalized.encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", ascii_text).strip().upper()


def parse_brazilian_number(token: str) -> float:
    cleaned = token.replace(".", "").replace(",", ".")
    return float(cleaned)


def compact_whitespace(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def unique_preserve_order(items: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for item in items:
        if item and item not in seen:
            seen.add(item)
            result.append(item)
    return result


def dedupe_consecutive(items: Iterable[str]) -> list[str]:
    result: list[str] = []
    previous = None
    for item in items:
        cleaned = compact_whitespace(item)
        if cleaned and cleaned != previous:
            result.append(cleaned)
            previous = cleaned
    return result


def extract_process_id(text: str) -> tuple[str, str]:
    cleaned = compact_whitespace(text)
    match = PROCESS_RE.search(cleaned)
    if not match:
        return "", cleaned
    process_id = match.group(1)
    remainder = compact_whitespace(f"{cleaned[: match.start()]} {cleaned[match.end() :]}")
    return process_id, remainder


def build_precatorio_number(process_id: str) -> str:
    if not re.fullmatch(r"\d{12}", process_id):
        return ""
    return f"{process_id[6:]}{process_id[2:4]}"


def canonical_payment_name(normalized_match: str) -> str:
    if normalized_match.startswith("PARCIAL"):
        return "Parcial Antecipação"
    if "PAGAMENTO ANTECIPACAO" in normalized_match:
        return "Pagamento Antecipação"
    if "PAGAMENTO INTEGRAL" in normalized_match:
        return "Pagamento Integral"
    if "DIRETO" in normalized_match:
        return "Cessão - Acordo Direto"
    return "Cessão - Acordo"


def extract_payment_types(chunks: Iterable[str]) -> list[str]:
    text = normalize_ascii("\n".join(chunks))
    return [canonical_payment_name(match.group(0)) for match in PAYMENT_RE.finditer(text)]


def infer_missing_payment_types(
    raw_chunks: Iterable[str],
    payment_types: list[str],
    dates: list[str],
    gross_tokens: list[str],
) -> list[str]:
    expected_components = min(len(dates), len(gross_tokens))
    missing_count = expected_components - len(payment_types)
    if missing_count <= 0:
        return payment_types

    raw_text = normalize_ascii("\n".join(raw_chunks))
    inferred_count = min(missing_count, len(ORPHAN_PAYMENT_RE.findall(raw_text)))
    if inferred_count <= 0:
        return payment_types

    inferred_type = canonical_payment_name("PAGAMENTO ANTECIPACAO")
    return payment_types + [inferred_type] * inferred_count


def trim_redundant_items(items: list[str], target_length: int) -> list[str]:
    trimmed = list(items)
    while len(trimmed) > target_length:
        removed = False
        for index in range(len(trimmed) - 1, -1, -1):
            if trimmed[index] in trimmed[:index]:
                trimmed.pop(index)
                removed = True
                break
        if not removed:
            trimmed.pop()
    return trimmed


def align_payment_sequences(
    payment_types: list[str],
    dates: list[str],
    gross_tokens: list[str],
) -> tuple[list[str], list[str], list[str]]:
    non_zero_counts = [count for count in (len(payment_types), len(dates), len(gross_tokens)) if count > 0]
    if not non_zero_counts:
        return payment_types, dates, gross_tokens

    target_length = min(non_zero_counts)
    payment_types = trim_redundant_items(payment_types, target_length)
    dates = trim_redundant_items(dates, target_length)
    gross_tokens = trim_redundant_items(gross_tokens, target_length)
    return payment_types, dates, gross_tokens


def assign_column(center_x: float) -> str | None:
    for column_name, (min_x, max_x) in COLUMN_RANGES.items():
        if min_x <= center_x < max_x:
            return column_name
    return None


def build_page_lines(page: fitz.Page) -> list[ExtractedLine]:
    transformed_words: list[tuple[float, float, float, float, str]] = []
    for x0, y0, x1, y1, text, *_ in page.get_text("words"):
        transformed = fitz.Rect(x0, y0, x1, y1) * page.derotation_matrix
        transformed_words.append((transformed.x0, transformed.y0, transformed.x1, transformed.y1, text))

    transformed_words.sort(key=lambda item: (-item[1], -item[0]))

    grouped_lines: list[tuple[float, list[tuple[float, float, float, float, str]]]] = []
    current_y: float | None = None
    current_words: list[tuple[float, float, float, float, str]] = []

    for word in transformed_words:
        x0, y0, x1, y1, text = word
        if current_y is None or abs(current_y - y0) <= LINE_Y_TOLERANCE:
            if current_y is None:
                current_y = y0
            current_words.append(word)
            continue

        grouped_lines.append((current_y, current_words))
        current_y = y0
        current_words = [word]

    if current_words:
        grouped_lines.append((current_y, current_words))

    extracted_lines: list[ExtractedLine] = []
    for y, words in grouped_lines:
        columns: dict[str, list[tuple[float, str]]] = {key: [] for key in COLUMN_RANGES}
        extras: list[tuple[float, str]] = []

        for x0, y0, x1, y1, text in words:
            column_name = assign_column((x0 + x1) / 2)
            if column_name is None:
                extras.append((x0, text))
                continue
            columns[column_name].append((x0, text))

        collapsed_columns = {
            key: compact_whitespace(" ".join(text for _, text in sorted(values, reverse=True)))
            for key, values in columns.items()
        }

        line_text = compact_whitespace(
            " ".join(
                filter(
                    None,
                    [*collapsed_columns.values(), " ".join(text for _, text in sorted(extras, reverse=True))],
                )
            )
        )

        extracted_lines.append(
            ExtractedLine(
                y=y,
                columns=collapsed_columns,
                text=line_text,
                normalized_text=normalize_ascii(line_text),
            )
        )

    return extracted_lines


def is_structural_line(line: ExtractedLine, top_limit: float) -> bool:
    if not line.text:
        return True
    if PAGE_RE.fullmatch(line.text):
        return True
    if line.y > top_limit:
        return True
    if line.normalized_text.startswith(SUMMARY_LINE_PREFIXES):
        return True
    if any(pattern in line.normalized_text for pattern in HEADER_PATTERNS):
        return True
    words = set(line.normalized_text.split())
    if len(words & HEADER_TOKEN_GUARD) >= 2 and not PROCESS_RE.search(line.normalized_text):
        return True
    return False


def build_name_segment(line: ExtractedLine) -> str:
    process_value = line.columns["processo"]
    creditor_value = line.columns["credor"]
    if not process_value and not creditor_value:
        return ""

    _, process_remainder = extract_process_id(process_value)
    segment_parts = []
    if process_remainder:
        segment_parts.append(process_remainder)
    if creditor_value:
        segment_parts.append(creditor_value)
    return compact_whitespace(" ".join(segment_parts))


def collect_record_bands(page: fitz.Page) -> list[dict[str, list[str]]]:
    page_lines = build_page_lines(page)
    process_lines = [line for line in page_lines if extract_process_id(line.columns["processo"])[0]]
    if not process_lines:
        return []

    top_limit = max(line.y for line in process_lines) + TOP_BAND_PADDING
    clean_lines = [line for line in page_lines if not is_structural_line(line, top_limit)]
    process_indices = [index for index, line in enumerate(clean_lines) if extract_process_id(line.columns["processo"])[0]]
    if not process_indices:
        return []

    process_ys = [clean_lines[index].y for index in process_indices]
    boundaries = [float("inf")]
    boundaries.extend((previous + current) / 2 for previous, current in zip(process_ys, process_ys[1:]))
    boundaries.append(float("-inf"))

    records: list[dict[str, list[str]]] = []
    for band_index in range(len(process_indices)):
        upper = boundaries[band_index]
        lower = boundaries[band_index + 1]
        record = {column_name: [] for column_name in COLUMN_RANGES}
        record[NAME_SEGMENTS_KEY] = []
        for line in clean_lines:
            if upper >= line.y > lower:
                for column_name, value in line.columns.items():
                    if value:
                        record[column_name].append(value)
                name_segment = build_name_segment(line)
                if name_segment:
                    record[NAME_SEGMENTS_KEY].append(name_segment)
        records.append(record)
    return records


def build_payment_components(
    record: dict[str, list[str]],
) -> tuple[list[PaymentComponent], list[str], PaymentBuildDiagnostics]:
    raw_payment_types = extract_payment_types(record["tipo_pagamento"])
    dates = DATE_RE.findall(" ".join(record["data_pagamento"]))
    gross_tokens = VALUE_RE.findall(" ".join(record["valor_bruto"]))
    payment_types = infer_missing_payment_types(record["tipo_pagamento"], raw_payment_types, dates, gross_tokens)
    inferred_payment_type_count = max(0, len(payment_types) - len(raw_payment_types))

    aligned_payment_types, aligned_dates, aligned_gross_tokens = align_payment_sequences(
        payment_types,
        dates,
        gross_tokens,
    )
    sequence_adjusted = (
        aligned_payment_types != payment_types
        or aligned_dates != dates
        or aligned_gross_tokens != gross_tokens
    )

    diagnostics = PaymentBuildDiagnostics(
        raw_payment_type_count=len(raw_payment_types),
        raw_date_count=len(dates),
        raw_gross_count=len(gross_tokens),
        inferred_payment_type_count=inferred_payment_type_count,
        sequence_adjusted=sequence_adjusted,
    )

    warnings: list[str] = []
    if not aligned_payment_types or not aligned_dates or not aligned_gross_tokens:
        warnings.append("Nao foi possivel alinhar componentes de pagamento.")
        return [], warnings, diagnostics

    if not (len(aligned_payment_types) == len(aligned_dates) == len(aligned_gross_tokens)):
        warnings.append("Quantidade divergente entre tipos, datas e valores brutos.")
        return [], warnings, diagnostics

    components = [
        PaymentComponent(
            payment_type=payment_type,
            date=date,
            gross_value=parse_brazilian_number(gross_token),
        )
        for payment_type, date, gross_token in zip(aligned_payment_types, aligned_dates, aligned_gross_tokens)
    ]
    return components, warnings, diagnostics


def reduce_payment_components(
    components: list[PaymentComponent],
) -> tuple[str, str, float | None, int, bool]:
    if not components:
        return "", "", None, 0, False

    retained = [component for component in components if not component.payment_type.startswith("Cessão")]
    had_cession = len(retained) != len(components)
    if not retained:
        retained = components

    type_text = " + ".join(unique_preserve_order(component.payment_type for component in retained))
    date_text = " | ".join(unique_preserve_order(component.date for component in retained))
    gross_total = round(sum(component.gross_value for component in retained), 2)
    return type_text, date_text, gross_total, len(retained), had_cession


def join_name_fragments(record: dict[str, list[str]], process_remainder: str) -> str:
    line_segments = dedupe_consecutive(record.get(NAME_SEGMENTS_KEY, []))
    if line_segments:
        return compact_whitespace(" ".join(line_segments))

    fragments = []
    if process_remainder:
        creditor_text = " ".join(record["credor"])
        normalized_remainder = normalize_ascii(process_remainder)
        normalized_creditor = normalize_ascii(creditor_text)
        if normalized_remainder and normalized_remainder not in normalized_creditor:
            fragments.append(process_remainder)
    fragments.extend(record["credor"])
    return compact_whitespace(" ".join(dedupe_consecutive(fragments)))


def join_entity_fragments(record: dict[str, list[str]]) -> str:
    cleaned_fragments = dedupe_consecutive(record["entidade"])
    return compact_whitespace(" ".join(unique_preserve_order(cleaned_fragments)))


def extract_liquid_value(record: dict[str, list[str]]) -> float | None:
    tokens = dedupe_consecutive(VALUE_RE.findall(" ".join(record["valor_liquido"])))
    if not tokens:
        return None
    return parse_brazilian_number(tokens[0])


def extract_process_ids_in_band(record: dict[str, list[str]]) -> list[str]:
    process_ids: list[str] = []
    for chunk in record["processo"]:
        process_id, _ = extract_process_id(chunk)
        if process_id and process_id not in process_ids:
            process_ids.append(process_id)
    return process_ids


def requires_manual_review(built_row: BuiltRow) -> bool:
    return bool(
        built_row.warnings
        or (MANUAL_REVIEW_FLAGS & set(built_row.diagnostic_flags))
    )


def has_auto_adjustment(built_row: BuiltRow) -> bool:
    return bool(AUTO_ADJUSTMENT_FLAGS & set(built_row.diagnostic_flags))


def count_manual_review_rows(built_rows: Iterable[BuiltRow]) -> int:
    return sum(1 for built_row in built_rows if requires_manual_review(built_row))


def count_auto_adjustment_rows(built_rows: Iterable[BuiltRow]) -> int:
    return sum(1 for built_row in built_rows if has_auto_adjustment(built_row))


def build_output_row(record: dict[str, list[str]], source_page: int = 0) -> BuiltRow:
    process_id, process_remainder = extract_process_id(" ".join(record["processo"]))
    components, warnings, diagnostics = build_payment_components(record)
    payment_type, payment_date, gross_total, retained_count, had_cession = reduce_payment_components(components)
    liquid_value = extract_liquid_value(record)
    process_ids_in_band = extract_process_ids_in_band(record)
    liquid_tokens = dedupe_consecutive(VALUE_RE.findall(" ".join(record["valor_liquido"])))

    row = {
        "Elaborador": "",
        "Nº do processo": process_id,
        "Nº precatorio": build_precatorio_number(process_id),
        "Nome do Credor": join_name_fragments(record, process_remainder),
        "Entidade/Ente Federado": join_entity_fragments(record),
        "Data pagamento": payment_date,
        "Tipo de pagamento": payment_type,
        "Valor Bruto do Pagamento": gross_total,
        "Valor líquido pago a parte": liquid_value,
        "NATUREZA": "",
    }

    if not process_id:
        warnings.append("Processo nao identificado.")
    if not payment_type:
        warnings.append("Tipo de pagamento nao identificado.")
    if gross_total is None:
        warnings.append("Valor bruto nao identificado.")

    diagnostic_flags: list[str] = []
    validation_notes: list[str] = []

    if diagnostics.inferred_payment_type_count:
        diagnostic_flags.append("inferred_payment_type")
        validation_notes.append(
            f"Tipo de pagamento inferido automaticamente em {diagnostics.inferred_payment_type_count} componente(s)."
        )
    if diagnostics.sequence_adjusted:
        diagnostic_flags.append("sequence_adjusted")
        validation_notes.append("Sequencia de tipos, datas ou valores ajustada automaticamente.")
    if len(liquid_tokens) > 1:
        diagnostic_flags.append("multiple_liquid_tokens")
        validation_notes.append(
            f"Faixa contem {len(liquid_tokens)} valores liquidos; o primeiro valor foi usado na planilha."
        )
    if len(process_ids_in_band) > 1:
        diagnostic_flags.append("multiple_process_ids")
        validation_notes.append(
            f"Faixa contem mais de um processo identificado: {', '.join(process_ids_in_band)}."
        )
    if gross_total is not None and liquid_value is not None and liquid_value > gross_total:
        diagnostic_flags.append("liquid_gt_gross")
        validation_notes.append("Valor liquido maior que o valor bruto consolidado.")

    return BuiltRow(
        row=row,
        component_count=len(components),
        retained_component_count=retained_count,
        had_cession=had_cession,
        warnings=warnings,
        diagnostic_flags=diagnostic_flags,
        validation_notes=validation_notes,
        source_page=source_page,
    )


def process_pdf(pdf_path: Path) -> list[BuiltRow]:
    built_rows: list[BuiltRow] = []
    with fitz.open(pdf_path) as document:
        for page_number, page in enumerate(document, start=1):
            for record in collect_record_bands(page):
                built_rows.append(build_output_row(record, source_page=page_number))
    return built_rows


def export_excel(rows: list[dict[str, object]], output_path: Path) -> None:
    dataframe = pd.DataFrame(rows, columns=OUTPUT_COLUMNS)
    dataframe.to_excel(output_path, index=False, sheet_name="Precatorios")

    workbook = openpyxl.load_workbook(output_path)
    worksheet = workbook["Precatorios"]

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    header_fill = PatternFill(fill_type="solid", fgColor="D9E2F3")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    column_widths = {
        "A": 16,
        "B": 18,
        "C": 16,
        "D": 44,
        "E": 28,
        "F": 22,
        "G": 34,
        "H": 24,
        "I": 24,
        "J": 18,
    }
    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

    currency_columns = {"H", "I"}
    text_columns = {"B", "C"}
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            if cell.column_letter in currency_columns and cell.value not in (None, ""):
                cell.number_format = 'R$ #,##0.00'
            if cell.column_letter in text_columns:
                cell.number_format = "@"

    workbook.save(output_path)


def format_brl(value: float | int | None) -> str:
    if value is None:
        return "n/d"
    formatted = f"{float(value):,.2f}"
    return "R$ " + formatted.replace(",", "_").replace(".", ",").replace("_", ".")


def format_delta(value: float | int) -> str:
    sign = "+" if value > 0 else ""
    formatted = f"{float(value):,.2f}"
    return sign + formatted.replace(",", "_").replace(".", ",").replace("_", ".")


def extract_pdf_declared_summary(pdf_path: Path) -> dict[str, float | int | None]:
    with fitz.open(pdf_path) as document:
        start_page = max(0, document.page_count - 3)
        text = "\n".join(document[index].get_text("text") for index in range(start_page, document.page_count))

    lines = [compact_whitespace(line) for line in text.splitlines() if compact_whitespace(line)]
    normalized_lines = [normalize_ascii(line) for line in lines]

    def find_money(label: str) -> float | None:
        for index, normalized_line in enumerate(normalized_lines):
            if label not in normalized_line:
                continue
            for offset in (1, -1, 2, -2, 3, -3):
                probe = index + offset
                if 0 <= probe < len(lines):
                    match = VALUE_RE.search(lines[probe])
                    if match:
                        return parse_brazilian_number(match.group(0))
        return None

    def find_integer(label: str) -> int | None:
        for index, normalized_line in enumerate(normalized_lines):
            if label not in normalized_line:
                continue
            for offset in (-1, 1, -2, 2, -3, 3):
                probe = index + offset
                if 0 <= probe < len(lines) and re.fullmatch(r"\d+", lines[probe]):
                    return int(lines[probe])
        return None

    return {
        "declared_record_count": find_integer("TOTAL DE REGISTROS"),
        "declared_party_gross_total": find_money("VALOR TOTAL DA PARTE APOS DESAGIO"),
        "declared_party_liquid_total": find_money("VALOR LIQUIDO PAGO A PARTE"),
    }


def aggregate_pdf_declared_summaries(pdf_files: list[Path]) -> tuple[dict[str, float | int | None], int]:
    summary = {
        "declared_record_count": 0,
        "declared_party_gross_total": 0.0,
        "declared_party_liquid_total": 0.0,
    }
    parsed_count = 0

    for pdf_file in pdf_files:
        current = extract_pdf_declared_summary(pdf_file)
        if any(current.values()):
            parsed_count += 1
        if current["declared_record_count"] is not None:
            summary["declared_record_count"] += int(current["declared_record_count"])
        if current["declared_party_gross_total"] is not None:
            summary["declared_party_gross_total"] += float(current["declared_party_gross_total"])
        if current["declared_party_liquid_total"] is not None:
            summary["declared_party_liquid_total"] += float(current["declared_party_liquid_total"])

    if parsed_count == 0:
        return {
            "declared_record_count": None,
            "declared_party_gross_total": None,
            "declared_party_liquid_total": None,
        }, 0

    return summary, parsed_count


def write_validation_report(report_path: Path, built_rows: list[BuiltRow], pdf_files: list[Path]) -> None:
    total_rows = len(built_rows)
    rows_with_cession = sum(1 for built_row in built_rows if built_row.had_cession)
    rows_with_multiple_kept_payments = sum(
        1 for built_row in built_rows if built_row.retained_component_count > 1
    )
    rows_with_inferred_payment_type = sum(
        1 for built_row in built_rows if "inferred_payment_type" in built_row.diagnostic_flags
    )
    rows_with_sequence_adjustment = sum(
        1 for built_row in built_rows if "sequence_adjusted" in built_row.diagnostic_flags
    )
    rows_with_multiple_liquid_tokens = sum(
        1 for built_row in built_rows if "multiple_liquid_tokens" in built_row.diagnostic_flags
    )
    rows_with_multiple_process_ids = sum(
        1 for built_row in built_rows if "multiple_process_ids" in built_row.diagnostic_flags
    )
    rows_with_liquid_gt_gross = sum(
        1 for built_row in built_rows if "liquid_gt_gross" in built_row.diagnostic_flags
    )

    extracted_gross_total = round(
        sum((built_row.row["Valor Bruto do Pagamento"] or 0) for built_row in built_rows),
        2,
    )
    extracted_liquid_total = round(
        sum((built_row.row["Valor líquido pago a parte"] or 0) for built_row in built_rows),
        2,
    )

    declared_summary, parsed_summary_count = aggregate_pdf_declared_summaries(pdf_files)
    declared_record_count = declared_summary["declared_record_count"]
    declared_gross_total = declared_summary["declared_party_gross_total"]
    declared_liquid_total = declared_summary["declared_party_liquid_total"]

    manual_review_rows = [built_row for built_row in built_rows if requires_manual_review(built_row)]
    auto_adjustment_rows = [
        built_row
        for built_row in built_rows
        if has_auto_adjustment(built_row) and not requires_manual_review(built_row)
    ]

    lines = [
        f"PDFs processados: {len(pdf_files)}",
        f"Arquivos: {', '.join(pdf.name for pdf in pdf_files)}",
        f"Registros extraidos: {total_rows}",
        f"Valor bruto consolidado extraido: {format_brl(extracted_gross_total)}",
        f"Valor liquido consolidado extraido: {format_brl(extracted_liquid_total)}",
        f"Registros com cessao filtrada: {rows_with_cession}",
        f"Registros com soma de multiplos pagamentos mantidos: {rows_with_multiple_kept_payments}",
        f"Registros com tipo inferido automaticamente: {rows_with_inferred_payment_type}",
        f"Registros com ajuste automatico de sequencia: {rows_with_sequence_adjustment}",
        f"Registros com mais de um valor liquido na faixa: {rows_with_multiple_liquid_tokens}",
        f"Faixas com mais de um processo identificado: {rows_with_multiple_process_ids}",
        f"Registros com liquido maior que bruto: {rows_with_liquid_gt_gross}",
        f"Registros com ajuste automatico monitorado: {count_auto_adjustment_rows(built_rows)}",
        f"Pendencias de validacao: {sum(1 for built_row in built_rows if built_row.warnings)}",
        f"Casos para revisao manual: {len(manual_review_rows)}",
    ]

    if parsed_summary_count:
        gross_delta = round(
            extracted_gross_total - float(declared_gross_total or 0),
            2,
        )
        liquid_delta = round(
            extracted_liquid_total - float(declared_liquid_total or 0),
            2,
        )
        record_delta = total_rows - int(declared_record_count or 0)

        lines.extend(
            [
                "",
                f"Resumo oficial identificado em: {parsed_summary_count}/{len(pdf_files)} PDF(s)",
                f"Total de registros do PDF: {declared_record_count}",
                f"Diferenca de registros extraidos x PDF: {record_delta}",
                f"Valor total da parte apos desagio no PDF: {format_brl(float(declared_gross_total or 0))}",
                f"Diferenca do bruto extraido x PDF: {format_delta(gross_delta)}",
                f"Valor liquido pago a parte no PDF: {format_brl(float(declared_liquid_total or 0))}",
                f"Diferenca do liquido extraido x PDF: {format_delta(liquid_delta)}",
            ]
        )
    else:
        lines.extend(
            [
                "",
                "Resumo oficial do PDF: nao localizado nas ultimas paginas processadas.",
            ]
        )

    if manual_review_rows:
        lines.extend(["", "Casos para revisao manual:"])
        for built_row in manual_review_rows:
            notes = built_row.validation_notes + built_row.warnings
            lines.append(
                f"p.{built_row.source_page} | {built_row.row['Nº do processo']} | "
                f"{built_row.row['Nome do Credor']} | {'; '.join(notes)}"
            )

    if auto_adjustment_rows:
        lines.extend(["", "Ajustes automaticos monitorados:"])
        for built_row in auto_adjustment_rows:
            notes = built_row.validation_notes + built_row.warnings
            lines.append(
                f"p.{built_row.source_page} | {built_row.row['Nº do processo']} | "
                f"{built_row.row['Nome do Credor']} | {'; '.join(notes)}"
            )

    report_path.write_text("\n".join(lines), encoding="utf-8")


def discover_pdfs(input_path: Path) -> list[Path]:
    if input_path.is_file():
        return [input_path]
    return sorted(path for path in input_path.glob("*.pdf") if path.is_file())


def run_extraction(input_path: Path, output_path: Path, report_path: Path) -> ExtractionRunResult:
    pdf_files = discover_pdfs(input_path)
    if not pdf_files:
        raise FileNotFoundError(f"Nenhum PDF encontrado em: {input_path}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.parent.mkdir(parents=True, exist_ok=True)

    built_rows: list[BuiltRow] = []
    for pdf_file in pdf_files:
        built_rows.extend(process_pdf(pdf_file))

    rows = [built_row.row for built_row in built_rows]
    export_excel(rows, output_path)
    write_validation_report(report_path, built_rows, pdf_files)
    return ExtractionRunResult(
        pdf_files=pdf_files,
        built_rows=built_rows,
        output_path=output_path,
        report_path=report_path,
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Extrai pagamentos de precatorios do PDF e gera uma planilha Excel consolidada."
    )
    parser.add_argument(
        "--input",
        default="entrada",
        help="Arquivo PDF ou pasta com PDFs. Padrao: entrada",
    )
    parser.add_argument(
        "--output",
        default="saida/precatorios_extraidos.xlsx",
        help="Caminho do arquivo Excel de saida. Padrao: saida/precatorios_extraidos.xlsx",
    )
    parser.add_argument(
        "--report",
        default="saida/precatorios_validacao.txt",
        help="Caminho do relatorio de validacao. Padrao: saida/precatorios_validacao.txt",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    input_path = Path(args.input).resolve()
    output_path = Path(args.output).resolve()
    report_path = Path(args.report).resolve()
    try:
        result = run_extraction(input_path, output_path, report_path)
    except FileNotFoundError as exc:
        print(str(exc))
        return 1

    unresolved_warnings = sum(1 for built_row in result.built_rows if built_row.warnings)
    manual_review_count = count_manual_review_rows(result.built_rows)
    auto_adjustment_count = count_auto_adjustment_rows(result.built_rows)
    print(f"PDFs processados: {len(result.pdf_files)}")
    print(f"Registros extraidos: {len(result.built_rows)}")
    print(f"Pendencias de validacao: {unresolved_warnings}")
    print(f"Casos para revisao manual: {manual_review_count}")
    print(f"Ajustes automaticos monitorados: {auto_adjustment_count}")
    print(f"Planilha gerada em: {result.output_path}")
    print(f"Relatorio gerado em: {result.report_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
