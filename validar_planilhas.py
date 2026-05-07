from __future__ import annotations

import argparse
import math
import re
import unicodedata
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill

from leitor_pdf import OUTPUT_COLUMNS, process_pdf


@dataclass(frozen=True)
class SheetMapping:
    process_col: str
    precatorio_col: str | None
    name_col: str
    entity_col: str | None
    date_col: str | None
    type_col: str | None
    gross_col: str
    liquid_col: str | None
    nature_col: str | None


@dataclass
class ComparableRecord:
    source: str
    row_number: int
    processo: str
    precatorio: str
    credor: str
    credor_match: str
    ente: str
    ente_norm: str
    data_raw: str
    data_norm: str
    tipo_raw: str
    tipo_norm: str
    valor_bruto: float | None
    valor_liquido: float | None
    natureza_raw: str
    natureza_norm: str
    sheet_name: str

    def exact_key(self) -> tuple[str, str, float | None, float | None]:
        return (self.processo, self.credor_match, self.valor_bruto, self.valor_liquido)

    def financial_key(self) -> tuple[str, float | None, float | None]:
        return (self.processo, self.valor_bruto, self.valor_liquido)


@dataclass
class ValidationRunResult:
    manual_workbook: Path
    sheet_name: str
    summary: dict[str, int]
    output_path: Path
    report_path: Path


def normalize_ascii(text: str) -> str:
    normalized = unicodedata.normalize("NFKD", text)
    ascii_text = normalized.encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", ascii_text).strip().upper()


def normalize_name_for_match(value: object) -> str:
    if pd.isna(value):
        return ""
    text = normalize_ascii(str(value))
    text = re.sub(r"[^A-Z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def normalize_text(value: object) -> str:
    if pd.isna(value):
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize_header(value: object) -> str:
    return normalize_ascii(str(value))


def normalize_process(value: object) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            if math.isnan(value):
                return ""
        except TypeError:
            pass
        return str(int(round(float(value))))
    return re.sub(r"\D", "", str(value))


def normalize_money(value: object) -> float | None:
    if pd.isna(value):
        return None
    if isinstance(value, str):
        text = value.replace("R$", "").replace(".", "").replace(",", ".")
        text = re.sub(r"[^0-9.-]", "", text.strip())
        if not text:
            return None
        return round(float(text), 2)
    return round(float(value), 2)


def normalize_date_token(token: str) -> str:
    token = token.strip()
    if not token:
        return ""
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(token, fmt).strftime("%d/%m/%Y")
        except ValueError:
            continue
    parsed = pd.to_datetime(token, errors="coerce", dayfirst=True)
    if pd.notna(parsed):
        return parsed.strftime("%d/%m/%Y")
    parsed = pd.to_datetime(token, errors="coerce", dayfirst=False)
    if pd.notna(parsed):
        return parsed.strftime("%d/%m/%Y")
    return token


def unique_preserve_order(items: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for item in items:
        if item and item not in seen:
            seen.add(item)
            result.append(item)
    return result


def normalize_date_field(value: object) -> str:
    if pd.isna(value):
        return ""
    raw_text = str(value).replace(" 00:00:00", "")
    tokens = [normalize_date_token(token) for token in raw_text.splitlines()]
    return " | ".join(unique_preserve_order(token for token in tokens if token))


def detect_sheet_mapping(sheet_name: str, dataframe: pd.DataFrame) -> SheetMapping | None:
    normalized_columns = {normalize_header(column): column for column in dataframe.columns}
    normalized_sheet = normalize_header(sheet_name)

    if normalized_sheet == "PARA ESTORNAR" and len(dataframe.columns) >= 11:
        return SheetMapping(
            process_col=dataframe.columns[1],
            precatorio_col=dataframe.columns[3],
            name_col=dataframe.columns[4],
            entity_col=dataframe.columns[5],
            date_col=dataframe.columns[6],
            type_col=dataframe.columns[7],
            gross_col=dataframe.columns[8],
            liquid_col=dataframe.columns[9],
            nature_col=dataframe.columns[10],
        )

    if all(key in normalized_columns for key in ("NO DO PROCESSO", "NOME DO CREDOR", "VALOR BRUTO DO PAGAMENTO")):
        return SheetMapping(
            process_col=normalized_columns["NO DO PROCESSO"],
            precatorio_col=normalized_columns.get("NUMERO DO PRECATORIO") or normalized_columns.get("NO PRECATORIO"),
            name_col=normalized_columns["NOME DO CREDOR"],
            entity_col=normalized_columns.get("ENTIDADE/ENTE FEDERADO"),
            date_col=normalized_columns.get("DATA PAGAMENTO"),
            type_col=normalized_columns.get("TIPO DE PAGAMENTO"),
            gross_col=normalized_columns["VALOR BRUTO DO PAGAMENTO"],
            liquid_col=normalized_columns.get("VALOR LIQUIDO PAGO A PARTE"),
            nature_col=normalized_columns.get("NATUREZA"),
        )

    has_process = "NO DO PROCESSO" in normalized_columns
    has_name = "NOME DO CREDOR" in normalized_columns
    has_layout_fields = all(
        key in normalized_columns
        for key in ("DATA PAGAMENTO", "TIPO DE PAGAMENTO", "VALOR LIQUIDO PAGO A PARTE")
    )
    if has_process and has_name and has_layout_fields and len(dataframe.columns) >= 10:
        return SheetMapping(
            process_col=dataframe.columns[1],
            precatorio_col=dataframe.columns[3] if len(dataframe.columns) > 3 else None,
            name_col=dataframe.columns[4],
            entity_col=dataframe.columns[5] if len(dataframe.columns) > 5 else None,
            date_col=dataframe.columns[6] if len(dataframe.columns) > 6 else None,
            type_col=dataframe.columns[7] if len(dataframe.columns) > 7 else None,
            gross_col=dataframe.columns[8],
            liquid_col=dataframe.columns[9] if len(dataframe.columns) > 9 else None,
            nature_col=dataframe.columns[10] if len(dataframe.columns) > 10 else None,
        )

    return None


def build_records(
    dataframe: pd.DataFrame,
    mapping: SheetMapping,
    source: str,
    sheet_name: str,
) -> list[ComparableRecord]:
    records: list[ComparableRecord] = []
    for row_index, row in dataframe.iterrows():
        processo = normalize_process(row[mapping.process_col])
        credor = normalize_text(row[mapping.name_col])
        if not processo and not credor:
            continue

        record = ComparableRecord(
            source=source,
            row_number=row_index + 2,
            processo=processo,
            precatorio=normalize_process(row[mapping.precatorio_col]) if mapping.precatorio_col else "",
            credor=credor,
            credor_match=normalize_name_for_match(row[mapping.name_col]),
            ente=normalize_text(row[mapping.entity_col]) if mapping.entity_col else "",
            ente_norm=normalize_ascii(normalize_text(row[mapping.entity_col])) if mapping.entity_col else "",
            data_raw=normalize_text(row[mapping.date_col]) if mapping.date_col else "",
            data_norm=normalize_date_field(row[mapping.date_col]) if mapping.date_col else "",
            tipo_raw=normalize_text(row[mapping.type_col]) if mapping.type_col else "",
            tipo_norm=normalize_ascii(normalize_text(row[mapping.type_col])) if mapping.type_col else "",
            valor_bruto=normalize_money(row[mapping.gross_col]),
            valor_liquido=normalize_money(row[mapping.liquid_col]) if mapping.liquid_col else None,
            natureza_raw=normalize_text(row[mapping.nature_col]) if mapping.nature_col else "",
            natureza_norm=normalize_ascii(normalize_text(row[mapping.nature_col])) if mapping.nature_col else "",
            sheet_name=sheet_name,
        )
        records.append(record)
    return records


def choose_manual_sheet(workbook_path: Path, auto_records: list[ComparableRecord], explicit_sheet: str | None) -> tuple[str, SheetMapping, list[ComparableRecord]]:
    workbook = pd.ExcelFile(workbook_path)
    auto_proc_name = {(record.processo, record.credor_match) for record in auto_records}

    if explicit_sheet:
        dataframe = workbook.parse(explicit_sheet)
        mapping = detect_sheet_mapping(explicit_sheet, dataframe)
        if mapping is None:
            raise ValueError(f"Nao foi possivel interpretar a aba informada: {explicit_sheet}")
        records = build_records(dataframe, mapping, "manual", explicit_sheet)
        return explicit_sheet, mapping, records

    best_sheet: str | None = None
    best_mapping: SheetMapping | None = None
    best_records: list[ComparableRecord] = []
    best_overlap = -1

    for sheet_name in workbook.sheet_names:
        dataframe = workbook.parse(sheet_name)
        mapping = detect_sheet_mapping(sheet_name, dataframe)
        if mapping is None:
            continue
        records = build_records(dataframe, mapping, "manual", sheet_name)
        if not records:
            continue
        overlap = len(auto_proc_name & {(record.processo, record.credor_match) for record in records})
        if overlap > best_overlap:
            best_sheet = sheet_name
            best_mapping = mapping
            best_records = records
            best_overlap = overlap

    if best_sheet is None or best_mapping is None:
        raise ValueError("Nenhuma aba comparavel foi encontrada na planilha manual.")

    return best_sheet, best_mapping, best_records


def compare_optional_field(manual_value: str, auto_value: str) -> bool:
    return bool(manual_value and auto_value and manual_value != auto_value)


def describe_pair_differences(manual_record: ComparableRecord, auto_record: ComparableRecord) -> str:
    differences: list[str] = []
    if compare_optional_field(manual_record.ente_norm, auto_record.ente_norm):
        differences.append("ente")
    if compare_optional_field(manual_record.data_norm, auto_record.data_norm):
        differences.append("data")
    if compare_optional_field(manual_record.tipo_norm, auto_record.tipo_norm):
        differences.append("tipo")
    return ", ".join(differences)


def pair_records(
    manual_records: list[ComparableRecord],
    auto_records: list[ComparableRecord],
) -> tuple[
    list[tuple[ComparableRecord, ComparableRecord]],
    list[tuple[ComparableRecord, ComparableRecord]],
    list[tuple[ComparableRecord, ComparableRecord]],
    list[ComparableRecord],
    list[ComparableRecord],
]:
    manual_by_exact: dict[tuple[str, str, float | None, float | None], list[ComparableRecord]] = defaultdict(list)
    auto_by_exact: dict[tuple[str, str, float | None, float | None], list[ComparableRecord]] = defaultdict(list)

    for record in manual_records:
        manual_by_exact[record.exact_key()].append(record)
    for record in auto_records:
        auto_by_exact[record.exact_key()].append(record)

    exact_matches: list[tuple[ComparableRecord, ComparableRecord]] = []
    remaining_manual: list[ComparableRecord] = []
    remaining_auto: list[ComparableRecord] = []

    for key in set(manual_by_exact) | set(auto_by_exact):
        manual_items = manual_by_exact.get(key, [])
        auto_items = auto_by_exact.get(key, [])
        match_count = min(len(manual_items), len(auto_items))
        exact_matches.extend(zip(manual_items[:match_count], auto_items[:match_count]))
        remaining_manual.extend(manual_items[match_count:])
        remaining_auto.extend(auto_items[match_count:])

    manual_by_financial: dict[tuple[str, float | None, float | None], list[ComparableRecord]] = defaultdict(list)
    auto_by_financial: dict[tuple[str, float | None, float | None], list[ComparableRecord]] = defaultdict(list)

    for record in remaining_manual:
        manual_by_financial[record.financial_key()].append(record)
    for record in remaining_auto:
        auto_by_financial[record.financial_key()].append(record)

    financial_matches: list[tuple[ComparableRecord, ComparableRecord]] = []
    remaining_manual_after_financial: list[ComparableRecord] = []
    remaining_auto_after_financial: list[ComparableRecord] = []

    for key in set(manual_by_financial) | set(auto_by_financial):
        manual_items = manual_by_financial.get(key, [])
        auto_items = auto_by_financial.get(key, [])
        match_count = min(len(manual_items), len(auto_items))
        financial_matches.extend(zip(manual_items[:match_count], auto_items[:match_count]))
        remaining_manual_after_financial.extend(manual_items[match_count:])
        remaining_auto_after_financial.extend(auto_items[match_count:])

    manual_by_name: dict[tuple[str, str], list[ComparableRecord]] = defaultdict(list)
    auto_by_name: dict[tuple[str, str], list[ComparableRecord]] = defaultdict(list)

    for record in remaining_manual_after_financial:
        manual_by_name[(record.processo, record.credor_match)].append(record)
    for record in remaining_auto_after_financial:
        auto_by_name[(record.processo, record.credor_match)].append(record)

    name_matches_value_diff: list[tuple[ComparableRecord, ComparableRecord]] = []
    only_manual: list[ComparableRecord] = []
    only_auto: list[ComparableRecord] = []

    for key in set(manual_by_name) | set(auto_by_name):
        manual_items = manual_by_name.get(key, [])
        auto_items = auto_by_name.get(key, [])
        match_count = min(len(manual_items), len(auto_items))
        name_matches_value_diff.extend(zip(manual_items[:match_count], auto_items[:match_count]))
        only_manual.extend(manual_items[match_count:])
        only_auto.extend(auto_items[match_count:])

    return exact_matches, financial_matches, name_matches_value_diff, only_manual, only_auto


def build_pdf_confirmation_set(pdf_path: Path) -> set[tuple[str, str, float | None, float | None]]:
    confirmation_set: set[tuple[str, str, float | None, float | None]] = set()
    for built_row in process_pdf(pdf_path):
        row = built_row.row
        confirmation_set.add(
            (
                normalize_process(row[OUTPUT_COLUMNS[1]]),
                normalize_name_for_match(row[OUTPUT_COLUMNS[3]]),
                normalize_money(row[OUTPUT_COLUMNS[7]]),
                normalize_money(row[OUTPUT_COLUMNS[8]]),
            )
        )
    return confirmation_set


def make_result_row(
    status: str,
    manual_record: ComparableRecord | None,
    auto_record: ComparableRecord | None,
    pdf_confirmed: bool | None,
    note: str,
) -> dict[str, object]:
    return {
        "status": status,
        "confirmado_no_pdf": pdf_confirmed,
        "processo": manual_record.processo if manual_record else auto_record.processo,
        "credor_manual": manual_record.credor if manual_record else "",
        "credor_automatico": auto_record.credor if auto_record else "",
        "valor_bruto_manual": manual_record.valor_bruto if manual_record else None,
        "valor_bruto_automatico": auto_record.valor_bruto if auto_record else None,
        "valor_liquido_manual": manual_record.valor_liquido if manual_record else None,
        "valor_liquido_automatico": auto_record.valor_liquido if auto_record else None,
        "data_manual": manual_record.data_norm if manual_record else "",
        "data_automatica": auto_record.data_norm if auto_record else "",
        "tipo_manual": manual_record.tipo_raw if manual_record else "",
        "tipo_automatico": auto_record.tipo_raw if auto_record else "",
        "ente_manual": manual_record.ente if manual_record else "",
        "ente_automatico": auto_record.ente if auto_record else "",
        "precatorio_manual": manual_record.precatorio if manual_record else "",
        "precatorio_automatico": auto_record.precatorio if auto_record else "",
        "natureza_manual": manual_record.natureza_raw if manual_record else "",
        "natureza_automatica": auto_record.natureza_raw if auto_record else "",
        "linha_manual": manual_record.row_number if manual_record else None,
        "linha_automatica": auto_record.row_number if auto_record else None,
        "aba_manual": manual_record.sheet_name if manual_record else "",
        "observacao": note,
    }


def write_excel_report(output_path: Path, summary_rows: list[dict[str, object]], comparison_rows: list[dict[str, object]]) -> None:
    summary_df = pd.DataFrame(summary_rows)
    comparison_df = pd.DataFrame(comparison_rows)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary_df.to_excel(writer, index=False, sheet_name="Resumo")
        comparison_df.to_excel(writer, index=False, sheet_name="Comparacao")
        comparison_df[comparison_df["status"] == "somente_automatica"].to_excel(
            writer,
            index=False,
            sheet_name="Somente_Automatica",
        )
        comparison_df[comparison_df["status"] == "somente_manual"].to_excel(
            writer,
            index=False,
            sheet_name="Somente_Manual",
        )
        comparison_df[comparison_df["status"] == "match_financeiro_nome_diferente"].to_excel(
            writer,
            index=False,
            sheet_name="Nome_Diferente",
        )
        comparison_df[comparison_df["status"] == "match_processo_nome_valor_diferente"].to_excel(
            writer,
            index=False,
            sheet_name="Valor_Diferente",
        )

    workbook = openpyxl.load_workbook(output_path)
    header_fill = PatternFill(fill_type="solid", fgColor="D9E2F3")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for column_cells in sheet.columns:
            max_width = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_width + 2, 14), 48)
    workbook.save(output_path)


def write_text_report(
    report_path: Path,
    workbook_path: Path,
    sheet_name: str,
    auto_path: Path,
    summary: dict[str, object],
    comparison_rows: list[dict[str, object]],
) -> None:
    lines = [
        f"Planilha manual: {workbook_path.name}",
        f"Aba utilizada: {sheet_name}",
        f"Planilha automatica: {auto_path.name}",
        f"Registros manuais comparados: {summary['registros_manuais']}",
        f"Registros automaticos comparados: {summary['registros_automaticos']}",
        f"Matches exatos: {summary['matches_exatos']}",
        f"Matches financeiros com nome diferente: {summary['matches_financeiros_nome_diferente']}",
        f"Matches por processo+nome com valor diferente: {summary['matches_processo_nome_valor_diferente']}",
        f"Somente manual: {summary['somente_manual']}",
        f"Somente automatica: {summary['somente_automatica']}",
        f"Somente automatica confirmada no PDF: {summary['somente_automatica_confirmada_pdf']}",
    ]

    name_diff_rows = [row for row in comparison_rows if row["status"] == "match_financeiro_nome_diferente"]
    value_diff_rows = [row for row in comparison_rows if row["status"] == "match_processo_nome_valor_diferente"]
    only_auto_rows = [row for row in comparison_rows if row["status"] == "somente_automatica"]

    if name_diff_rows:
        lines.append("")
        lines.append("Divergencias de nome com mesmos valores:")
        for row in name_diff_rows:
            lines.append(
                f"{row['processo']} | manual: {row['credor_manual']} | automatico: {row['credor_automatico']} | bruto: {row['valor_bruto_automatico']} | liquido: {row['valor_liquido_automatico']}"
            )

    if only_auto_rows:
        lines.append("")
        lines.append("Registros presentes so na saida automatica:")
        for row in only_auto_rows:
            confirmado = "sim" if row["confirmado_no_pdf"] else "nao"
            lines.append(
                f"{row['processo']} | {row['credor_automatico']} | bruto: {row['valor_bruto_automatico']} | liquido: {row['valor_liquido_automatico']} | confirmado no PDF: {confirmado}"
            )

    if value_diff_rows:
        lines.append("")
        lines.append("Registros com mesmo processo e credor, mas valores ou metadados diferentes:")
        for row in value_diff_rows:
            lines.append(
                f"{row['processo']} | {row['credor_automatico'] or row['credor_manual']} | bruto manual: {row['valor_bruto_manual']} | bruto automatico: {row['valor_bruto_automatico']} | liquido manual: {row['valor_liquido_manual']} | liquido automatico: {row['valor_liquido_automatico']}"
            )

    report_path.write_text("\n".join(lines), encoding="utf-8")


def resolve_manual_workbook(manual_source: Path) -> Path:
    if manual_source.is_file():
        return manual_source

    manual_candidates = sorted(path for path in manual_source.glob("*.xlsx") if not path.name.startswith("~$"))
    if not manual_candidates:
        raise FileNotFoundError(f"Nenhuma planilha manual encontrada em: {manual_source}")
    return manual_candidates[0]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Compara a planilha manual da pasta de validacao com a planilha extraida automaticamente."
    )
    parser.add_argument(
        "--manual",
        "--manual-dir",
        dest="manual",
        default="validação",
        help="Pasta onde esta a planilha manual. Padrao: validação",
    )
    parser.add_argument(
        "--auto",
        default="saida/precatorios_extraidos.xlsx",
        help="Planilha automatica gerada pelo extrator.",
    )
    parser.add_argument(
        "--pdf",
        default="entrada/Relatorio_Precatorios.pdf",
        help="PDF original para confirmar registros somente automaticos.",
    )
    parser.add_argument(
        "--sheet",
        default=None,
        help="Nome da aba manual a ser usada. Se omitido, o script escolhe a melhor aba automaticamente.",
    )
    parser.add_argument(
        "--output",
        default="saida/validacao_comparativa.xlsx",
        help="Arquivo Excel de saida do comparativo.",
    )
    parser.add_argument(
        "--report",
        default="saida/validacao_comparativa.txt",
        help="Arquivo texto com o resumo do comparativo.",
    )
    return parser.parse_args()


def run_validation(
    manual_source: Path,
    auto_path: Path,
    pdf_path: Path,
    output_path: Path,
    report_path: Path,
    sheet: str | None = None,
) -> ValidationRunResult:
    manual_workbook = resolve_manual_workbook(manual_source)
    auto_df = pd.read_excel(auto_path)
    auto_mapping = SheetMapping(
        process_col=auto_df.columns[1],
        precatorio_col=auto_df.columns[2],
        name_col=auto_df.columns[3],
        entity_col=auto_df.columns[4],
        date_col=auto_df.columns[5],
        type_col=auto_df.columns[6],
        gross_col=auto_df.columns[7],
        liquid_col=auto_df.columns[8],
        nature_col=auto_df.columns[9],
    )
    auto_records = build_records(auto_df, auto_mapping, "automatica", "Precatorios")

    sheet_name, _, manual_records = choose_manual_sheet(manual_workbook, auto_records, sheet)
    exact_matches, financial_matches, process_name_matches, only_manual, only_auto = pair_records(manual_records, auto_records)

    pdf_confirmation = build_pdf_confirmation_set(pdf_path) if pdf_path.exists() else set()

    comparison_rows: list[dict[str, object]] = []
    for manual_record, auto_record in exact_matches:
        note = describe_pair_differences(manual_record, auto_record)
        comparison_rows.append(
            make_result_row(
                "match_exato",
                manual_record,
                auto_record,
                None,
                note,
            )
        )

    for manual_record, auto_record in financial_matches:
        note = "Mesmo processo e mesmos valores, mas com divergencia de nome."
        comparison_rows.append(
            make_result_row(
                "match_financeiro_nome_diferente",
                manual_record,
                auto_record,
                None,
                note,
            )
        )

    for manual_record, auto_record in process_name_matches:
        note = "Mesmo processo e mesmo credor, mas com divergencia de valores ou campos complementares."
        comparison_rows.append(
            make_result_row(
                "match_processo_nome_valor_diferente",
                manual_record,
                auto_record,
                None,
                note,
            )
        )

    for manual_record in only_manual:
        comparison_rows.append(
            make_result_row(
                "somente_manual",
                manual_record,
                None,
                None,
                "Existe na planilha manual e nao foi encontrado na saida automatica.",
            )
        )

    for auto_record in only_auto:
        confirmed_in_pdf = auto_record.exact_key() in pdf_confirmation
        note = "Existe na saida automatica e foi confirmado no PDF original." if confirmed_in_pdf else "Existe na saida automatica, mas nao foi confirmado pelo PDF."
        comparison_rows.append(
            make_result_row(
                "somente_automatica",
                None,
                auto_record,
                confirmed_in_pdf,
                note,
            )
        )

    comparison_rows.sort(
        key=lambda row: (
            row["status"],
            row["processo"] or "",
            row["credor_automatico"] or row["credor_manual"] or "",
        )
    )

    summary = {
        "registros_manuais": len(manual_records),
        "registros_automaticos": len(auto_records),
        "matches_exatos": len(exact_matches),
        "matches_financeiros_nome_diferente": len(financial_matches),
        "matches_processo_nome_valor_diferente": len(process_name_matches),
        "somente_manual": len(only_manual),
        "somente_automatica": len(only_auto),
        "somente_automatica_confirmada_pdf": sum(1 for row in comparison_rows if row["status"] == "somente_automatica" and row["confirmado_no_pdf"]),
    }

    summary_rows = [
        {"metrica": "planilha_manual", "valor": manual_workbook.name},
        {"metrica": "aba_utilizada", "valor": sheet_name},
        {"metrica": "planilha_automatica", "valor": auto_path.name},
        {"metrica": "registros_manuais", "valor": summary["registros_manuais"]},
        {"metrica": "registros_automaticos", "valor": summary["registros_automaticos"]},
        {"metrica": "matches_exatos", "valor": summary["matches_exatos"]},
        {
            "metrica": "matches_financeiros_nome_diferente",
            "valor": summary["matches_financeiros_nome_diferente"],
        },
        {
            "metrica": "matches_processo_nome_valor_diferente",
            "valor": summary["matches_processo_nome_valor_diferente"],
        },
        {"metrica": "somente_manual", "valor": summary["somente_manual"]},
        {"metrica": "somente_automatica", "valor": summary["somente_automatica"]},
        {
            "metrica": "somente_automatica_confirmada_pdf",
            "valor": summary["somente_automatica_confirmada_pdf"],
        },
    ]

    output_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    write_excel_report(output_path, summary_rows, comparison_rows)
    write_text_report(report_path, manual_workbook, sheet_name, auto_path, summary, comparison_rows)

    return ValidationRunResult(
        manual_workbook=manual_workbook,
        sheet_name=sheet_name,
        summary=summary,
        output_path=output_path,
        report_path=report_path,
    )


def main() -> int:
    args = parse_args()
    manual_source = Path(args.manual).resolve()
    auto_path = Path(args.auto).resolve()
    pdf_path = Path(args.pdf).resolve()
    output_path = Path(args.output).resolve()
    report_path = Path(args.report).resolve()

    try:
        result = run_validation(
            manual_source=manual_source,
            auto_path=auto_path,
            pdf_path=pdf_path,
            output_path=output_path,
            report_path=report_path,
            sheet=args.sheet,
        )
    except FileNotFoundError as exc:
        print(str(exc))
        return 1

    summary = result.summary
    print(f"Aba manual utilizada: {result.sheet_name}")
    print(f"Registros manuais: {summary['registros_manuais']}")
    print(f"Registros automaticos: {summary['registros_automaticos']}")
    print(f"Matches exatos: {summary['matches_exatos']}")
    print(f"Matches financeiros com nome diferente: {summary['matches_financeiros_nome_diferente']}")
    print(f"Matches por processo+nome com valor diferente: {summary['matches_processo_nome_valor_diferente']}")
    print(f"Somente manual: {summary['somente_manual']}")
    print(f"Somente automatica: {summary['somente_automatica']}")
    print(f"Somente automatica confirmada no PDF: {summary['somente_automatica_confirmada_pdf']}")
    print(f"Excel de validacao: {result.output_path}")
    print(f"Resumo de validacao: {result.report_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
